"""Build the neighborhood intake workbook: every field that populates WordPress."""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

NEIGHBORHOODS = ['Art Museum Area', 'Avenue of the Arts', 'Bella Vista', 'Chinatown',
                 'Fairmount', 'Fishtown', 'Fitler Square', 'Graduate Hospital',
                 'Logan Square', 'Northern Liberties', 'Old City', 'Queen Village',
                 'Rittenhouse Square', 'Society Hill', 'Washington Square',
                 'Washington Square West']
YESNO = ['Yes', 'No']

FONT = 'Arial'
HEAD_BG = PatternFill('solid', fgColor='2B2B2B')
GROUP_BG = PatternFill('solid', fgColor='83838B')
SEED_BG = PatternFill('solid', fgColor='F4F2EF')
THIN = Side(style='thin', color='D9D9D9')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# (group, header, width, wrap)
COLUMNS = [
    ('Record',    'Neighborhood Name', 24, False),
    ('Record',    'URL Slug', 22, False),
    ('Record',    'Published', 12, False),

    ('Card',      'Card Image', 26, False),
    ('Card',      'Median Price', 16, False),
    ('Card',      'Destination URL', 30, False),
    ('Card',      'Featured', 11, False),
    ('Card',      'Featured Rank', 13, False),

    ('Hero',      'Hero Image', 26, False),
    ('Hero',      'Hero Descriptor', 46, True),

    ('Content',   'Neighborhood Overview', 60, True),
    ('Content',   'Neighborhood Real Estate', 60, True),
    ('Content',   'Real Estate Image', 26, False),
    ('Content',   "Andy's Perspective — Heading", 32, True),
    ('Content',   "Andy's Perspective", 60, True),

    ('Related',   'Featured Condominium Buildings', 40, True),
    ('Related',   'IDX Listings Embed', 34, True),

    ('Closing',   'Closing Statement', 46, True),

    ('Internal',  'ZIP Codes', 18, False),
    ('Internal',  'Philadelphia Section', 22, False),
    ('Internal',  'Internal Research Record', 50, True),
]

GUIDANCE = {
    'Neighborhood Name': 'Required. As it appears on the card and page.',
    'URL Slug': 'Leave blank to use the name. Existing 16 are already set.',
    'Published': 'Yes or No. No keeps it a draft.',
    'Card Image': 'Filename of the image you upload to the Media Library.',
    'Median Price': 'Free text. "Data pending" is fine.',
    'Destination URL': 'Only if the card should link somewhere other than this page.',
    'Featured': 'Yes or No. Featured ones lead the directory.',
    'Featured Rank': 'Number. 1 shows first. Only for featured rows.',
    'Hero Image': 'Filename of the image you upload to the Media Library.',
    'Hero Descriptor': '15–25 words.',
    'Neighborhood Overview': '100–125 words. One concise overview, not multiple lifestyle sections.',
    'Neighborhood Real Estate': '110–140 words. Buyer-first and condominium-led; avoid tourism copy.',
    'Real Estate Image': 'Filename of the image you upload to the Media Library.',
    "Andy's Perspective — Heading": 'Short heading specific to this neighborhood.',
    "Andy's Perspective": 'Judgment-led professional insight, specific to this neighborhood.',
    'Featured Condominium Buildings': 'Building names, separated by commas. Leave empty where condominium inventory is not meaningful.',
    'IDX Listings Embed': 'The IDX shortcode or embed code, filtered to this neighborhood.',
    'Closing Statement': 'Short closing line before the Explore and Inquire buttons.',
    'ZIP Codes': 'Internal. Comma-separated.',
    'Philadelphia Section': 'Internal. e.g. Center City, River Wards.',
    'Internal Research Record': 'Internal. Sources, methodology, maintenance notes. Never published.',
}

wb = Workbook()

# ── Read Me ───────────────────────────────────────────────────────────
rm = wb.active
rm.title = 'Read Me'
rm.sheet_view.showGridLines = False
rm.column_dimensions['A'].width = 112

LINES = [
    ('Neighborhoods — Intake', 16, True),
    ('', 11, False),
    ('Every field that populates a neighborhood in WordPress. One row per neighborhood,', 11, False),
    ('sixteen in all. All sixteen names are already filled in — write alongside them.', 11, False),
    ('', 11, False),
    ('Only Neighborhood Name is required. Everything else can be left blank and added later;', 11, False),
    ('a blank shows as a dash or is simply omitted, and nothing breaks.', 11, False),
    ('', 11, False),
    ('Where each column lands in WordPress', 13, True),
    ('', 11, False),
    ('  Record      The post itself — title, URL, published or draft.', 11, False),
    ('  Card        What shows on the Neighborhoods directory card.', 11, False),
    ('  Hero        Top of the individual neighborhood page.', 11, False),
    ('  Content     The written body of the page.', 11, False),
    ('  Related     Featured buildings and the live IDX listings.', 11, False),
    ('  Closing     The closing statement above the call to action.', 11, False),
    ('  Internal    Never published. Research and maintenance notes.', 11, False),
    ('', 11, False),
    ('Images', 13, True),
    ('', 11, False),
    ('Write the filename only, e.g. rittenhouse-square-hero.jpg. Upload the files to the', 11, False),
    ('Media Library separately; the names here just say which image goes where.', 11, False),
    ('', 11, False),
    ('FAQs', 13, True),
    ('', 11, False),
    ('FAQs are on their own tab, because each neighborhood can have several. Add one row', 11, False),
    ('per question and pick the neighborhood it belongs to from the dropdown.', 11, False),
    ('', 11, False),
    ('Send this back when the rows are in and it gets loaded into the site.', 11, False),
]

for i, (text, size, bold) in enumerate(LINES, start=1):
    c = rm.cell(row=i, column=1, value=text)
    c.font = Font(name=FONT, size=size, bold=bold, color='2B2B2B')

# ── Neighborhoods ─────────────────────────────────────────────────────
ws = wb.create_sheet('Neighborhoods')
ws.sheet_view.showGridLines = False
ws.freeze_panes = 'B4'

ws.row_dimensions[1].height = 18
ws.row_dimensions[2].height = 30
ws.row_dimensions[3].height = 56

for n, (group, head, width, wrap) in enumerate(COLUMNS, start=1):
    letter = get_column_letter(n)
    ws.column_dimensions[letter].width = width

    g = ws.cell(row=1, column=n, value=group)
    g.font = Font(name=FONT, size=9, bold=True, color='FFFFFF')
    g.fill = GROUP_BG
    g.alignment = Alignment(horizontal='center', vertical='center')

    h = ws.cell(row=2, column=n, value=head)
    h.font = Font(name=FONT, size=10, bold=True, color='FFFFFF')
    h.fill = HEAD_BG
    h.alignment = Alignment(vertical='center', wrap_text=True)

    note = ws.cell(row=3, column=n, value=GUIDANCE.get(head, ''))
    note.font = Font(name=FONT, size=9, italic=True, color='83838B')
    note.alignment = Alignment(vertical='top', wrap_text=True)

ROWS = len(NEIGHBORHOODS) + 6  # room for additions

for i in range(ROWS):
    row = 4 + i
    ws.row_dimensions[row].height = 46

    for n, (group, head, width, wrap) in enumerate(COLUMNS, start=1):
        c = ws.cell(row=row, column=n)
        c.font = Font(name=FONT, size=10, color='2B2B2B')
        c.border = BORDER
        c.alignment = Alignment(vertical='top', wrap_text=wrap)

    if i < len(NEIGHBORHOODS):
        name = ws.cell(row=row, column=1, value=NEIGHBORHOODS[i])
        name.font = Font(name=FONT, size=10, bold=True, color='2B2B2B')
        name.fill = SEED_BG
        name.alignment = Alignment(vertical='top')

        slug = NEIGHBORHOODS[i].lower().replace(' ', '-')
        s = ws.cell(row=row, column=2, value=slug)
        s.font = Font(name=FONT, size=10, color='83838B')
        s.fill = SEED_BG
        s.alignment = Alignment(vertical='top')

        p = ws.cell(row=row, column=3, value='Yes')
        p.font = Font(name=FONT, size=10, color='2B2B2B')
        p.alignment = Alignment(vertical='top')

# ── FAQs ──────────────────────────────────────────────────────────────
fq = wb.create_sheet('FAQs')
fq.sheet_view.showGridLines = False
fq.freeze_panes = 'A3'

FAQ_COLS = [('Neighborhood', 24), ('Question', 55), ('Answer', 85)]
FAQ_NOTES = ['Pick from the list.',
             'Factual, evergreen buyer question.',
             'Concise. Use the internal research record to keep it useful.']

for n, (head, width) in enumerate(FAQ_COLS, start=1):
    fq.column_dimensions[get_column_letter(n)].width = width

    h = fq.cell(row=1, column=n, value=head)
    h.font = Font(name=FONT, size=10, bold=True, color='FFFFFF')
    h.fill = HEAD_BG
    h.alignment = Alignment(vertical='center')

    note = fq.cell(row=2, column=n, value=FAQ_NOTES[n - 1])
    note.font = Font(name=FONT, size=9, italic=True, color='83838B')

fq.row_dimensions[1].height = 20

for row in range(3, 83):
    fq.row_dimensions[row].height = 40
    for n in range(1, 4):
        c = fq.cell(row=row, column=n)
        c.font = Font(name=FONT, size=10, color='2B2B2B')
        c.border = BORDER
        c.alignment = Alignment(vertical='top', wrap_text=(n == 3))

# ── Lists ─────────────────────────────────────────────────────────────
ls = wb.create_sheet('Lists')
ls.sheet_view.showGridLines = False
ls.column_dimensions['A'].width = 28
ls.column_dimensions['B'].width = 12

for n, (head, values) in enumerate([('Neighborhood', NEIGHBORHOODS), ('Yes / No', YESNO)], start=1):
    h = ls.cell(row=1, column=n, value=head)
    h.font = Font(name=FONT, size=10, bold=True, color='FFFFFF')
    h.fill = HEAD_BG

    for i, v in enumerate(values, start=2):
        c = ls.cell(row=i, column=n, value=v)
        c.font = Font(name=FONT, size=10, color='2B2B2B')

note = ls.cell(row=len(NEIGHBORHOODS) + 4, column=1,
               value='These are the values the site already knows. Adding one here does not '
                     'create it on the site — ask first if a new neighborhood is needed.')
note.font = Font(name=FONT, size=10, italic=True, color='83838B')


def dropdown( sheet, target_range, source, count, message ):
    dv = DataValidation(
        type='list',
        formula1='Lists!${0}$2:${0}${1}'.format( source, count + 1 ),
        allow_blank=True,
        showDropDown=False,  # False means "show the arrow" in the file format.
    )
    dv.showErrorMessage = True
    dv.errorTitle = 'Not on the list'
    dv.error = message
    sheet.add_data_validation( dv )
    dv.add( target_range )


dropdown( ws, 'C4:C{}'.format( ROWS + 3 ), 'B', len( YESNO ), 'Yes or No.' )
dropdown( ws, 'G4:G{}'.format( ROWS + 3 ), 'B', len( YESNO ), 'Yes or No.' )
dropdown( fq, 'A3:A82', 'A', len( NEIGHBORHOODS ),
          'Pick a neighborhood from the dropdown so the FAQ lands on the right page.' )

wb.save('/home/user/Brandon-/data/Andy-Oei-Neighborhoods-Intake.xlsx')
print('written')
